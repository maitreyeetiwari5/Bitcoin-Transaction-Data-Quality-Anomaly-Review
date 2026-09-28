"""Write the Excel exception report and CSV extracts for Tableau.

The Summary sheet uses live formulas that point at the detail sheets, so every headline
number can be traced back to the rows behind it.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
NOTE_FONT = Font(name=FONT, italic=True, color="595959")


def _write_table(ws, df: pd.DataFrame, start_row: int = 1):
    for c, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=str(name))
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(row, start=1):
            if isinstance(value, pd.Timestamp):
                value = value.tz_localize(None) if value.tzinfo else value
            try:
                if pd.isna(value):
                    value = None
            except (TypeError, ValueError):
                pass
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
    for c, name in enumerate(df.columns, start=1):
        sample = [len(str(name))] + [len(str(v)) for v in df.iloc[:200, c - 1]]
        ws.column_dimensions[get_column_letter(c)].width = min(max(sample) + 2, 60)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _summary(ws, has_json: bool):
    ws["A1"] = "Bitcoin Transaction Data Quality & Anomaly Review"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Every figure below is a live formula pointing at the detail sheets. "
                "Nothing on this sheet is typed in by hand.")
    ws["A2"].font = NOTE_FONT

    ev = "'Model Evaluation'"
    lines = [
        ("Data quality", None, None),
        ("Checks run", "=COUNTA('Data Quality Log'!A:A)-1", "Data Quality Log"),
        ("Checks with issues", "=COUNTIF('Data Quality Log'!D:D,\"ISSUE\")", "Data Quality Log"),
        ("Records affected by issues", "=SUM('Data Quality Log'!E:E)", "Data Quality Log"),
        ("SQL tie-outs matching pandas",
         "=COUNTIF('SQL Tie-Out'!E:E,\"MATCH\")&\" of \"&(COUNTA('SQL Tie-Out'!A:A)-1)",
         "SQL Tie-Out"),
        ("Anomaly review: holdout period (time steps 35-49)", None, None),
        ("Labeled transactions evaluated",
         f"=INDEX({ev}!B:B,MATCH(\"Labeled transactions evaluated\",{ev}!A:A,0))",
         "Model Evaluation"),
        ("Precision", f"=INDEX({ev}!B:B,MATCH(\"Precision\",{ev}!A:A,0))", "Model Evaluation"),
        ("Baseline illicit rate",
         f"=INDEX({ev}!B:B,MATCH(\"Baseline illicit rate\",{ev}!A:A,0))", "Model Evaluation"),
        ("Recall", f"=INDEX({ev}!B:B,MATCH(\"Recall\",{ev}!A:A,0))", "Model Evaluation"),
        ("Lift over baseline (x)",
         f"=INDEX({ev}!B:B,MATCH(\"Lift over baseline\",{ev}!A:A,0))", "Model Evaluation"),
        ("Precision in training period, for comparison",
         f"=INDEX({ev}!C:C,MATCH(\"Precision\",{ev}!A:A,0))", "Model Evaluation"),
        ("Unlabeled transactions escalated for review", "=COUNTA('Review Queue'!A:A)-1",
         "Review Queue"),
    ]
    if has_json:
        js = "'JSON Summary'"
        lines += [
            ("Raw JSON (semi-structured) check", None, None),
            ("Transactions parsed from JSON",
             f"=INDEX({js}!B:B,MATCH(\"Transactions parsed\",{js}!A:A,0))", "JSON Summary"),
            ("Fee reconciliation breaks",
             f"=INDEX({js}!B:B,MATCH(\"Fee reconciliation: breaks\",{js}!A:A,0))",
             "JSON Summary"),
        ]

    ws["A4"], ws["B4"], ws["C4"] = "Metric", "Value", "Source sheet"
    for cell in (ws["A4"], ws["B4"], ws["C4"]):
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    row = 5
    for label, formula, source in lines:
        ws.cell(row=row, column=1, value=label).font = (
            Font(name=FONT, bold=True) if formula is None else BODY_FONT)
        if formula:
            cell = ws.cell(row=row, column=2, value=formula)
            cell.font = BODY_FONT
            if label.startswith(("Precision", "Recall", "Baseline")):
                cell.number_format = "0.0%"
            elif label.startswith("Lift"):
                cell.number_format = '0.0"x"'
            ws.cell(row=row, column=3, value=source).font = NOTE_FONT
        row += 1
    ws.cell(row=row + 1, column=1,
            value="Unlabeled flags are escalated for analyst review, not classified as illicit."
            ).font = NOTE_FONT
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 20


def write_report(sheets: dict, json_tables: dict | None, path=None):
    path = path or config.OUTPUT_DIR / "exception_report.xlsx"
    wb = Workbook()
    _summary(wb.active, has_json=bool(json_tables))
    wb.active.title = "Summary"
    for name, df in {**sheets, **(json_tables or {})}.items():
        _write_table(wb.create_sheet(name), df)
    wb.save(path)
    return path


def export_csvs(scored: pd.DataFrame, by_step: pd.DataFrame, issue_log: pd.DataFrame):
    """Flat extracts sized for Tableau Public."""
    out = config.OUTPUT_DIR / "tableau"
    out.mkdir(parents=True, exist_ok=True)
    cols = ["txId", "time_step", "period", "label", "in_degree", "out_degree", "anomaly_score",
            "top_driver", "features_breached", "flagged"]
    scored[cols].to_csv(out / "transactions_scored.csv", index=False)
    by_step.to_csv(out / "evaluation_by_time_step.csv", index=False)
    issue_log.to_csv(out / "data_quality_log.csv", index=False)
    return out
